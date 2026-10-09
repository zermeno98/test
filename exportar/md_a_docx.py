"""Convierte los documentos Markdown del proyecto a Word (.docx).

Usa un .docx existente como plantilla: conserva sus estilos (Heading 1-4, Body Text,
First Paragraph, Compact, Block Text, Table, Hyperlink) y su encabezado (en 02 y 03,
el encabezado del formato F21-COOPYD-01). Agrega pie de página con número de página.

Uso:
    python md_a_docx.py ENTRADA.md PLANTILLA.docx SALIDA.docx [--estandar] [--confidencial TEXTO]
                        [--diagramas CARPETA_PNG] [--titulo-pie TEXTO]

    --estandar        omite las líneas iniciales hasta el primer '---' (las lleva el encabezado F21)
    --confidencial    agrega un encabezado con ese texto (para documentos confidenciales)
    --diagramas       sustituye en la exposición de motivos los esquemas de texto por los diagramas PNG
"""
import argparse
import copy
import re
from pathlib import Path

import docx
import markdown
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Inches, Pt, RGBColor
from lxml import html as LH

CARTA_ANCHO, CARTA_ALTO = 12240, 15840  # DXA


# ----------------------------------------------------------------------------- Markdown
def preprocesar(md: str) -> str:
    """Ajusta el Markdown para Python-Markdown.

    1. Las listas anidadas necesitan 4 espacios; los documentos usan 3.
    2. Una lista necesita una línea en blanco antes cuando sigue a un párrafo
       (por ejemplo, "**Apoyos/Requerimientos**" seguido de "- Equipo…").
    """
    patron_lista = re.compile(r'^\s*([-*+]|\d+\.) ')
    salida = []
    for linea in md.split('\n'):
        if re.match(r'^ {3}(- |\d+\. )', linea) and not linea.startswith('    '):
            linea = ' ' + linea
        if patron_lista.match(linea) and not linea.startswith(' ') and salida:
            previa = salida[-1]
            if (previa.strip() and not patron_lista.match(previa) and not previa.startswith(' ')
                    and not previa.lstrip().startswith(('|', '>'))):
                salida.append('')
        salida.append(linea)
    return '\n'.join(salida)


FIGURAS = {
    'FIGURA1': ('01-estandar-ciclo.png', 'Figura 1. Un estándar se crea una vez y certifica a muchas personas. Seis pasos en dos etapas.'),
    'FIGURA2': ('02-familia-estandares.png', 'Figura 2. Dos estándares: uno para quien usa la IA y otro para quien responde por ella. Familia de dos estándares y siete elementos.'),
    'FIGURA3': ('03-ruta-fases.png', 'Figura 3. Ruta de implementación: cuatro fases y tres puntos de control.'),
}


def sustituir_figuras(texto: str) -> str:
    """Mismos reemplazos que el PDF con diagramas (diagramas/generador/build_pdf.py)."""
    pares = [
        (r"\*\*Así nace un estándar\*\*.*?vuelve a evaluarse\.\n", "\n\nFIGURA1\n"),
        (r"\| EC-A · Usuario de IA \| EC-B · Responsable de la IA \|\n.*?incidentes e informe a la dirección \|\n", "\nFIGURA2\n"),
        (r"\| Fase \| Qué incluye \| Punto de control \|\n.*?revisión cada 2 años \| \|\n", "\nFIGURA3\n"),
    ]
    for patron, marcador in pares:
        texto, n = re.subn(patron, marcador, texto, flags=re.S)
        assert n == 1, marcador
    return texto


# ----------------------------------------------------------------------------- utilidades XML
def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    for w in tcPr.findall(qn('w:tcW')):
        tcPr.remove(w)
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(dxa)))
    tcW.set(qn('w:type'), 'dxa')
    tcPr.insert(0, tcW)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)


def table_borders(table, color='BFC0C0', sz='4'):
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    tblPr.append(borders)


def table_width(table, dxa, widths):
    tblPr = table._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')):
        tblPr.remove(old)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(dxa)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = table._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn('w:gridCol'))):
        gc.set(qn('w:w'), str(int(widths[i])))



ORDEN_TBLPR = ['tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize',
               'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook',
               'tblCaption', 'tblDescription']


def ordenar_tblpr(table):
    tblPr = table._tbl.tblPr
    hijos = list(tblPr)
    for h in hijos:
        tblPr.remove(h)
    clave = lambda e: ORDEN_TBLPR.index(e.tag.split('}')[1]) if e.tag.split('}')[1] in ORDEN_TBLPR else 99
    for h in sorted(hijos, key=clave):
        tblPr.append(h)

def margenes_celda(table, dxa):
    """Márgenes internos izquierdo y derecho de las celdas (tablas de muchas columnas)."""
    tblPr = table._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for lado in ('left', 'right'):
        el = OxmlElement(f'w:{lado}')
        el.set(qn('w:w'), str(dxa))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trPr.append(el)
    cant = OxmlElement('w:cantSplit')
    trPr.append(cant)


def add_field(run, instr):
    for kind, text in (('begin', None), ('instr', instr), ('separate', None), ('text', '1'), ('end', None)):
        if kind == 'instr':
            el = OxmlElement('w:instrText')
            el.set(qn('xml:space'), 'preserve')
            el.text = f' {text} '
            run._r.append(el)
        elif kind == 'text':
            t = OxmlElement('w:t')
            t.text = text
            run._r.append(t)
        else:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            run._r.append(fc)


def paragraph_bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), '4')
    b.set(qn('w:space'), '1')
    b.set(qn('w:color'), 'BFC0C0')
    pbdr.append(b)
    pPr.append(pbdr)



ESTILOS = {'Heading 1': 'Heading1', 'Heading 2': 'Heading2', 'Heading 3': 'Heading3', 'Heading 4': 'Heading4',
           'Body Text': 'BodyText', 'First Paragraph': 'FirstParagraph', 'Compact': 'Compact',
           'Block Text': 'BlockText', 'Image Caption': 'ImageCaption', 'Table': 'Table'}


def estilo(doc, nombre):
    """Busca el estilo por su identificador (los nombres de pandoc no coinciden con los de python-docx)."""
    sid = ESTILOS.get(nombre, nombre.replace(' ', ''))
    for s in doc.styles:
        if s.style_id == sid:
            return s
    raise KeyError(nombre)

def anchos_columnas(filas, ncols, total):
    """Anchos de columna (DXA) que no parten palabras.

    Cada columna recibe al menos el ancho de su palabra más larga; el espacio que sobra se
    reparte según cuánto texto tiene cada columna. Si la tabla no cabe con letra de 8.5 pt,
    se intenta con 7.5 pt. Devuelve (anchos, tamaño de letra).
    """
    angosta = ncols >= 6
    margen = 190 if angosta else 300  # márgenes internos de la celda y holgura
    datos = []
    for f in filas:
        celdas = f.findall('th') + f.findall('td')
        enc = bool(f.findall('th'))
        for i, c in enumerate(celdas[:ncols]):
            texto = ' '.join(c.text_content().split())
            palabras = texto.split(' ') if texto else ['']
            datos.append((i, enc, min(max(len(p) for p in palabras), 28), min(len(texto), 160)))
    for tam in ((8.5, 7.5, 7, 6.5) if angosta else (8.5, 7.5)):
        char = tam * 20 * 0.56          # ancho medio de un carácter de Arial, en DXA
        char_neg = tam * 20 * 0.62      # en negrita (encabezados)
        minimo = [700] * ncols
        natural = [700] * ncols
        for i, enc, larga, largo in datos:
            minimo[i] = max(minimo[i], larga * (char_neg if enc else char) + margen)
            natural[i] = max(natural[i], largo * (char_neg if enc else char) + margen)
        if sum(minimo) <= total:
            break
    if sum(minimo) > total:
        escala = total / sum(minimo)
        print(f'  aviso: tabla de {ncols} columnas no cabe sin partir palabras (escala {escala:.2f})')
        return [m * escala for m in minimo], tam
    sobrante = total - sum(minimo)
    necesidad = [max(n - m, 0) for n, m in zip(natural, minimo)]
    if sum(necesidad) <= sobrante:
        # todo cabe en una línea: repartir el resto en proporción al ancho natural
        anchos = [n for n in natural]
        resto = total - sum(anchos)
        anchos = [a + resto * a / sum(anchos) for a in anchos]
    else:
        anchos = [m + sobrante * nd / sum(necesidad) for m, nd in zip(minimo, necesidad)]
    return anchos, tam


# ----------------------------------------------------------------------------- convertidor
class Conversor:
    def __init__(self, doc, ancho_texto, carpeta_png=None):
        self.doc = doc
        self.ancho = ancho_texto
        self.png = carpeta_png
        self.prev_heading = True
        num_el = doc.part.numbering_part.element
        self.numbering = num_el
        ids = [int(n.get(qn('w:numId'))) for n in num_el.findall(qn('w:num'))]
        self.next_num = max(ids or [0]) + 1
        abstracts = {}
        for a in num_el.findall(qn('w:abstractNum')):
            fmt = a.find('.//' + qn('w:numFmt')).get(qn('w:val'))
            txt = a.find('.//' + qn('w:lvlText')).get(qn('w:val'))
            abstracts.setdefault(fmt, []).append((a.get(qn('w:abstractNumId')), txt))
        bullets = [i for i, t in abstracts.get('bullet', []) if t.strip()]
        self.abs_bullet = bullets[0] if bullets else self._crear_abstracto('8001', 'bullet')
        decimales = abstracts.get('decimal', [])
        self.abs_decimal = decimales[0][0] if decimales else self._crear_abstracto('8000', 'decimal')
        self.num_bullet = self._new_num(self.abs_bullet)

    def _crear_abstracto(self, abstract_id, formato):
        """Crea una definición de lista de 9 niveles cuando la plantilla no la trae."""
        a = OxmlElement('w:abstractNum')
        a.set(qn('w:abstractNumId'), abstract_id)
        mlt = OxmlElement('w:multiLevelType')
        mlt.set(qn('w:val'), 'multilevel')
        a.append(mlt)
        simbolos = ['•', '–', '•', '–', '•', '–', '•', '–', '•']
        for lvl in range(9):
            L = OxmlElement('w:lvl')
            L.set(qn('w:ilvl'), str(lvl))
            st = OxmlElement('w:start')
            st.set(qn('w:val'), '1')
            L.append(st)
            nf = OxmlElement('w:numFmt')
            if formato == 'decimal':
                nf.set(qn('w:val'), 'decimal' if lvl % 2 == 0 else 'lowerLetter')
            else:
                nf.set(qn('w:val'), 'bullet')
            L.append(nf)
            lt = OxmlElement('w:lvlText')
            lt.set(qn('w:val'), f'%{lvl + 1}.' if formato == 'decimal' else simbolos[lvl])
            L.append(lt)
            jc = OxmlElement('w:lvlJc')
            jc.set(qn('w:val'), 'left')
            L.append(jc)
            pPr = OxmlElement('w:pPr')
            ind = OxmlElement('w:ind')
            ind.set(qn('w:left'), str(720 + 360 * lvl))
            ind.set(qn('w:hanging'), '360')
            pPr.append(ind)
            L.append(pPr)
            a.append(L)
        primer_num = self.numbering.find(qn('w:num'))
        if primer_num is not None:
            primer_num.addprevious(a)
        else:
            self.numbering.append(a)
        return abstract_id

    def _new_num(self, abstract_id, start=None):
        num = OxmlElement('w:num')
        num.set(qn('w:numId'), str(self.next_num))
        a = OxmlElement('w:abstractNumId')
        a.set(qn('w:val'), str(abstract_id))
        num.append(a)
        if start is not None:
            for lvl in range(0, 9):
                ov = OxmlElement('w:lvlOverride')
                ov.set(qn('w:ilvl'), str(lvl))
                so = OxmlElement('w:startOverride')
                so.set(qn('w:val'), str(start if lvl == 0 else 1))
                ov.append(so)
                num.append(ov)
        self.numbering.append(num)
        self.next_num += 1
        return self.next_num - 1

    # --- inline
    def inline(self, par, el, bold=False, italic=False, code=False, size=None):
        if el.text:
            self._run(par, el.text, bold, italic, code, size)
        for child in el:
            tag = child.tag if isinstance(child.tag, str) else ''
            if tag in ('strong', 'b'):
                self.inline(par, child, True, italic, code, size)
            elif tag in ('em', 'i'):
                self.inline(par, child, bold, True, code, size)
            elif tag == 'code':
                self.inline(par, child, bold, italic, True, size)
            elif tag == 'a':
                self._link(par, child.text_content(), child.get('href', ''), bold, italic, size)
            elif tag == 'br':
                par.add_run().add_break(WD_BREAK.LINE)
            elif tag in ('ul', 'ol', 'p', 'table'):
                pass  # bloques dentro de un li: se procesan aparte
            else:
                self.inline(par, child, bold, italic, code, size)
            if child.tail:
                self._run(par, child.tail, bold, italic, code, size)

    def _run(self, par, text, bold, italic, code, size):
        text = text.replace('\n', ' ')
        if not text:
            return
        r = par.add_run(text)
        r.bold = bold or None
        r.italic = italic or None
        if code:
            r.font.name = 'Consolas'
        if size:
            r.font.size = Pt(size)

    def _link(self, par, text, url, bold, italic, size):
        if not url or url.startswith('#'):
            self._run(par, text, bold, italic, False, size)
            return
        r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement('w:hyperlink')
        h.set(qn('r:id'), r_id)
        r = OxmlElement('w:r')
        rPr = OxmlElement('w:rPr')
        rs = OxmlElement('w:rStyle')
        rs.set(qn('w:val'), 'Hyperlink')
        rPr.append(rs)
        if bold:
            rPr.append(OxmlElement('w:b'))
        if size:
            s = OxmlElement('w:sz')
            s.set(qn('w:val'), str(int(size * 2)))
            rPr.append(s)
        r.append(rPr)
        t = OxmlElement('w:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        r.append(t)
        h.append(r)
        par._p.append(h)

    # --- bloques
    def bloque(self, el, nivel_lista=0, estilo_p=None):
        tag = el.tag if isinstance(el.tag, str) else ''
        if tag in ('h1', 'h2', 'h3', 'h4', 'h5'):
            nivel = min(int(tag[1]), 4)
            p = self.doc.add_paragraph(style=estilo(self.doc, f'Heading {nivel}'))
            self.inline(p, el)
            self.prev_heading = True
            return
        if tag == 'p':
            txt = el.text_content().strip()
            if txt in FIGURAS and self.png:
                self.figura(txt)
                return
            est = estilo_p or ('First Paragraph' if self.prev_heading else 'Body Text')
            p = self.doc.add_paragraph(style=estilo(self.doc, est))
            self.inline(p, el)
            self.prev_heading = False
            return
        if tag in ('ul', 'ol'):
            self.lista(el, nivel_lista)
            self.prev_heading = False
            return
        if tag == 'table':
            self.tabla(el)
            self.prev_heading = False
            return
        if tag == 'blockquote':
            for child in el:
                self.bloque(child, nivel_lista, estilo_p='Block Text')
            self.prev_heading = False
            return
        if tag == 'hr':
            p = self.doc.add_paragraph(style=estilo(self.doc, 'Body Text'))
            paragraph_bottom_border(p)
            return
        if tag == 'pre':
            p = self.doc.add_paragraph(style=estilo(self.doc, 'Body Text'))
            self._run(p, el.text_content(), False, False, True, 9)
            return
        # otros: procesar hijos
        for child in el:
            self.bloque(child, nivel_lista, estilo_p)

    def lista(self, el, nivel):
        ordenada = el.tag == 'ol'
        if ordenada:
            start = int(el.get('start', '1'))
            num_id = self._new_num(self.abs_decimal, start)
        else:
            num_id = self.num_bullet
        for li in el.findall('li'):
            p = self.doc.add_paragraph(style=estilo(self.doc, 'Compact'))
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement('w:numPr')
            il = OxmlElement('w:ilvl')
            il.set(qn('w:val'), str(nivel))
            ni = OxmlElement('w:numId')
            ni.set(qn('w:val'), str(num_id))
            numPr.append(il)
            numPr.append(ni)
            pPr.append(numPr)
            # texto directo del li o del primer <p>
            primeros_p = [c for c in li if c.tag == 'p']
            if li.text and li.text.strip() or not primeros_p:
                self.inline(p, li)
            else:
                self.inline(p, primeros_p[0])
            for child in li:
                if child.tag in ('ul', 'ol'):
                    self.lista(child, nivel + 1)
                elif child.tag == 'p' and primeros_p and child is not primeros_p[0]:
                    q = self.doc.add_paragraph(style=estilo(self.doc, 'Compact'))
                    q.paragraph_format.left_indent = Inches(0.5 * (nivel + 1))
                    self.inline(q, child)
                elif child.tag == 'table':
                    self.tabla(child)

    def tabla(self, el):
        filas = el.findall('.//tr')
        if not filas:
            return
        ncols = max(len(f.findall('th') + f.findall('td')) for f in filas)
        t = self.doc.add_table(rows=0, cols=ncols)
        try:
            t.style = estilo(self.doc, 'Table')
        except KeyError:
            pass
        widths, tam = anchos_columnas(filas, ncols, self.ancho)
        if ncols >= 6:
            margenes_celda(t, 60)
        for fi, f in enumerate(filas):
            celdas = f.findall('th') + f.findall('td')
            row = t.add_row()
            es_enc = bool(f.findall('th'))
            if es_enc:
                repeat_header(row)
            for i in range(ncols):
                cell = row.cells[i]
                set_cell_width(cell, widths[i])
                p = cell.paragraphs[0]
                p.style = estilo(self.doc, 'Compact')
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                if i < len(celdas):
                    self.inline(p, celdas[i], bold=es_enc, size=tam)
                if es_enc:
                    shade(cell, 'ECECEC')
        table_borders(t)
        table_width(t, self.ancho, widths)
        ordenar_tblpr(t)
        self.doc.add_paragraph(style=estilo(self.doc, 'Body Text'))

    def figura(self, clave):
        archivo, pie = FIGURAS[clave]
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(Path(self.png) / archivo), width=Emu(int(self.ancho * 635)))
        try:
            c = self.doc.add_paragraph(pie, style=estilo(self.doc, 'Image Caption'))
        except KeyError:
            c = self.doc.add_paragraph(pie, style=estilo(self.doc, 'Body Text'))
        self.prev_heading = False


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('entrada')
    ap.add_argument('plantilla')
    ap.add_argument('salida')
    ap.add_argument('--estandar', action='store_true')
    ap.add_argument('--confidencial')
    ap.add_argument('--diagramas')
    ap.add_argument('--titulo-pie', default='')
    a = ap.parse_args()

    texto = Path(a.entrada).read_text(encoding='utf-8')
    if a.estandar:
        texto = texto.split('\n---\n', 1)[1]
    if a.diagramas:
        texto = sustituir_figuras(texto)
    html = markdown.markdown(preprocesar(texto), extensions=['tables', 'sane_lists'])
    raiz = LH.fragment_fromstring(html, create_parent='div')

    doc = docx.Document(a.plantilla)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    sec = doc.sections[0]
    sec.page_width = Emu(CARTA_ANCHO * 635)
    sec.page_height = Emu(CARTA_ALTO * 635)
    if a.estandar:
        izq = der = 1247
        sec.top_margin = Emu(1814 * 635)
    else:
        izq = der = 1304  # 0.9 in
        sec.top_margin = Emu(1304 * 635)
    sec.left_margin = Emu(izq * 635)
    sec.right_margin = Emu(der * 635)
    sec.bottom_margin = Emu(1134 * 635)
    sec.header_distance = Emu(567 * 635)
    sec.footer_distance = Emu(567 * 635)
    sec.gutter = Emu(0)
    ancho = CARTA_ANCHO - izq - der

    if a.confidencial:
        hp = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
        hp.text = ''
        r = hp.add_run(a.confidencial)
        r.bold = True
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # pie con número de página (salvo en los estándares, cuyo encabezado ya lo trae)
    if not a.estandar:
        fp = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        fp.text = ''
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = fp.add_run(f'{a.titulo_pie} · Página ' if a.titulo_pie else 'Página ')
        r.font.size = Pt(8)
        r2 = fp.add_run()
        r2.font.size = Pt(8)
        add_field(r2, 'PAGE')
        r3 = fp.add_run(' de ')
        r3.font.size = Pt(8)
        r4 = fp.add_run()
        r4.font.size = Pt(8)
        add_field(r4, 'NUMPAGES')

    conv = Conversor(doc, ancho, a.diagramas)
    for el in raiz:
        conv.bloque(el)
    # el sectPr debe quedar al final del cuerpo
    sect = body.find(qn('w:sectPr'))
    body.remove(sect)
    body.append(sect)

    h1 = raiz.find('h1')
    doc.core_properties.title = h1.text_content() if h1 is not None else Path(a.entrada).stem
    doc.core_properties.author = 'Proyecto Estándares de Competencia en IA'
    doc.core_properties.comments = 'Generado desde Markdown el 9 de octubre de 2026'
    doc.save(a.salida)
    print('ok', a.salida)


if __name__ == '__main__':
    main()
