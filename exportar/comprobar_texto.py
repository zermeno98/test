"""Comprueba que el texto del Markdown aparece completo en el .docx generado (palabras faltantes)."""
import re, sys, collections
import docx, markdown
from lxml import html as LH
sys.stdout.reconfigure(encoding='utf-8')
md, dx = sys.argv[1], sys.argv[2]
estandar = '--estandar' in sys.argv
texto = open(md, encoding='utf-8').read()
if estandar:
    texto = texto.split('\n---\n', 1)[1]
if '--diagramas' in sys.argv:
    # en la exposición de motivos, dos listas y dos tablas se sustituyen por diagramas
    sys.path.insert(0, __import__('os').path.dirname(__file__))
    from md_a_docx import sustituir_figuras
    texto = sustituir_figuras(texto)
sys.path.insert(0, __import__('os').path.dirname(__file__))
from md_a_docx import preprocesar
plano = LH.fragment_fromstring(markdown.markdown(preprocesar(texto), extensions=['tables', 'sane_lists']), create_parent='div').text_content()
d = docx.Document(dx)
partes = [p.text for p in d.paragraphs]
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            partes.append(c.text)
# hipervínculos: python-docx no incluye su texto en p.text en todas las versiones
xml = d.element.xml
partes += re.findall(r'<w:t[^>]*>([^<]*)</w:t>', xml)
dtxt = ' '.join(partes)
pal = lambda s: collections.Counter(re.findall(r'\w+', s.lower()))
a, b = pal(plano), pal(dtxt)
faltan = {w: n - b.get(w, 0) for w, n in a.items() if n > b.get(w, 0)}
print('palabras md', sum(a.values()), 'docx', sum(b.values()), 'faltantes', sum(faltan.values()), list(faltan.items())[:15])
